"""Liest den Feldkatalog (katalog/feldkatalog.xlsx).

Die Excel-Datei ist die maßgebliche Quelle für Felder, Bezeichnungen DE/EN, Datentypen,
Kernfelder, Auswahlwerte und das Wörterbuch. Fehlt sie im Programmordner, wird die
mitgelieferte Fassung dorthin kopiert.
"""
from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from . import pfade

# Datentypen (wie im Katalog)
TEXT = "Text"
ZAHL = "Zahl"
BEREICH = "Bereich (von–bis)"
JA_NEIN = "Ja/Nein"
AUSWAHL = "Auswahl"
MEHRFACH = "Mehrfachauswahl"
LISTE = "Liste"
SCHNITTSTELLE = "Schnittstelle"
DATENTYPEN = {TEXT, ZAHL, BEREICH, JA_NEIN, AUSWAHL, MEHRFACH, LISTE, SCHNITTSTELLE}

# Abschnitte
KOPF = "Kopfdaten"
ALLG = "Allgemeine Daten"
SCHN = "Schnittstellen"
TECH = "Technische Daten"
OPT = "Optionen"
LIEF = "Lieferumfang"
INTERN = "Intern (nicht auf Datenblatt)"

ABSCHNITT_EN = {ALLG: "General data", SCHN: "Interfaces", TECH: "Technical data", OPT: "Options",
                LIEF: "Scope of delivery", KOPF: "Header"}


class KatalogFehler(Exception):
    pass


@dataclass
class Feld:
    id: str
    abschnitt: str
    de: str
    en: str
    typ: str
    einheit: str = ""
    kern: bool = False
    filterart: str = ""
    pflicht: bool = False
    bemerkung: str = ""
    auswahl: list[tuple[str, str]] = field(default_factory=list)  # (de, en)

    def label(self, sprache: str) -> str:
        return self.en if sprache == "en" else self.de


@dataclass
class Katalog:
    felder: list[Feld]
    woerterbuch: dict[str, str]          # de -> en
    schnittstellen_werte: list[tuple[str, str]]
    quelle: Path
    synonyme: dict[str, str] = field(default_factory=dict)   # normalisierte Bezeichnung -> Feld-ID
    formulierungen: list = field(default_factory=list)        # [(englisch, deutsch)] – Textbausteine in Werten
    bezeichnungen: dict = field(default_factory=dict)         # (Gerätetyp, Feld-ID) -> (DE, EN)

    def bezeichnung(self, f: "Feld", geraetetypen: list[str], sprache: str) -> str:
        """Feldbezeichnung fürs Datenblatt; hat das Gerät genau einen Gerätetyp mit eigener Bezeichnung, gilt diese."""
        if len(geraetetypen) == 1 and (geraetetypen[0], f.id) in self.bezeichnungen:
            de, en = self.bezeichnungen[(geraetetypen[0], f.id)]
            return (en or de) if sprache == "en" else de
        return f.label(sprache)

    def __post_init__(self):
        self._by_id = {f.id: f for f in self.felder}

    def feld(self, fid: str) -> Feld:
        return self._by_id[fid]

    def hat(self, fid: str) -> bool:
        return fid in self._by_id

    def abschnitt(self, name: str) -> list[Feld]:
        return [f for f in self.felder if f.abschnitt == name]

    @property
    def kernfelder(self) -> list[Feld]:
        return [f for f in self.felder if f.kern]

    def uebersetze_rueck(self, text_en: str) -> str | None:
        """Feste Übersetzung Englisch → Deutsch (Wörterbuch rückwärts)."""
        if not text_en:
            return ""
        key = text_en.strip().lower()
        for de, en in self.woerterbuch.items():
            if en.strip().lower() == key:
                return de
        return None

    def teilweise(self, text: str, ziel: str) -> str:
        """Ersetzt bekannte Textbausteine („at full load“ ↔ „bei Volllast“). ziel = "de" oder "en".
        Zahlen und unbekannte Wörter bleiben stehen; das Ergebnis ist ein Vorschlag (gelb markieren)."""
        if not text or not self.formulierungen:
            return text
        paare = [(en, de) if ziel == "de" else (de, en) for en, de in self.formulierungen]
        paare = [(q, z) for q, z in paare if q and z and q.lower() != z.lower()]
        paare.sort(key=lambda p: len(p[0]), reverse=True)
        ersatz: list[str] = []
        t = text
        for quelle, zielwort in paare:
            muster = re.compile(r"(?<![\w])" + r"\s+".join(re.escape(w) for w in quelle.split()) + r"(?![\w])", re.I)

            def _sub(m, zielwort=zielwort):
                z = zielwort
                if m.group(0)[:1].isupper() and z[:1].islower():
                    z = z[:1].upper() + z[1:]
                ersatz.append(z)
                return f"\x00{len(ersatz) - 1}\x00"
            t = muster.sub(_sub, t)
        return re.sub(r"\x00(\d+)\x00", lambda m: ersatz[int(m.group(1))], t)

    def uebersetze(self, text_de: str) -> str | None:
        """Feste Übersetzung aus dem Wörterbuch (exakt, ohne Groß-/Kleinschreibung)."""
        if not text_de:
            return ""
        key = text_de.strip().lower()
        for de, en in self.woerterbuch.items():
            if de.strip().lower() == key:
                return en
        return None


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def sicherstellen() -> Path:
    """Kopiert den mitgelieferten Katalog in den Programmordner, falls dort keiner liegt."""
    ziel = pfade.KATALOG_DATEI
    if not ziel.exists() and pfade.KATALOG_MITGELIEFERT.exists() and pfade.KATALOG_MITGELIEFERT != ziel:
        ziel.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(pfade.KATALOG_MITGELIEFERT, ziel)
    return ziel


def laden(pfad: Path | None = None) -> Katalog:
    pfad = Path(pfad) if pfad else sicherstellen()
    if not pfad.exists():
        raise KatalogFehler(f"Feldkatalog nicht gefunden: {pfad}")
    try:
        wb = load_workbook(pfad, read_only=True, data_only=True)
    except PermissionError as e:
        raise KatalogFehler(f"Feldkatalog ist gesperrt (in Excel geöffnet?): {pfad}") from e
    for blatt in ("Felder", "Auswahlwerte", "Wörterbuch"):
        if blatt not in wb.sheetnames:
            raise KatalogFehler(f"Blatt „{blatt}“ fehlt im Feldkatalog.")

    rows = list(wb["Felder"].iter_rows(values_only=True))
    kopf = [_text(c) for c in rows[0]]
    idx = {name: kopf.index(name) for name in kopf if name}
    noetig = ["ID", "Abschnitt", "Feld (DE)", "Feld (EN)", "Datentyp"]
    for n in noetig:
        if n not in idx:
            raise KatalogFehler(f"Spalte „{n}“ fehlt im Blatt Felder.")

    def g(r, name):
        i = idx.get(name)
        return _text(r[i]) if i is not None and i < len(r) else ""

    felder: list[Feld] = []
    gesehen = set()
    for nr, r in enumerate(rows[1:], start=2):
        fid = g(r, "ID")
        if not fid:
            continue
        if fid in gesehen:
            raise KatalogFehler(f"ID „{fid}“ kommt doppelt vor (Zeile {nr}).")
        gesehen.add(fid)
        typ = g(r, "Datentyp")
        if typ not in DATENTYPEN:
            raise KatalogFehler(f"Unbekannter Datentyp „{typ}“ bei {fid} (Zeile {nr}).")
        felder.append(Feld(
            id=fid, abschnitt=g(r, "Abschnitt"), de=g(r, "Feld (DE)") or fid, en=g(r, "Feld (EN)") or g(r, "Feld (DE)"),
            typ=typ, einheit=g(r, "Einheit"), kern=g(r, "Kernfeld (18)").lower() == "ja",
            filterart=g(r, "Filterart Konfigurator"), pflicht=g(r, "Pflicht").lower() == "ja",
            bemerkung=g(r, "Bemerkung")))

    by_id = {f.id: f for f in felder}
    schnittstellen_werte: list[tuple[str, str]] = []
    for r in list(wb["Auswahlwerte"].iter_rows(values_only=True))[1:]:
        fid, de, en = _text(r[0]), _text(r[1]), _text(r[2] if len(r) > 2 else "")
        if not fid or not de:
            continue
        if fid.startswith("Schnittstelle"):
            schnittstellen_werte.append((de, en or de))
        elif fid in by_id:
            by_id[fid].auswahl.append((de, en or de))
    if not schnittstellen_werte:
        schnittstellen_werte = [("Standard", "Standard"), ("Option", "Option"), ("—", "—")]

    woerter: dict[str, str] = {}
    for r in list(wb["Wörterbuch"].iter_rows(values_only=True))[1:]:
        de, en = _text(r[0]), _text(r[1] if len(r) > 1 else "")
        if de and en:
            woerter[de] = en
    # Auswahlwerte gelten ebenfalls als feste Übersetzungen
    for f in felder:
        for de, en in f.auswahl:
            woerter.setdefault(de, en)
    for de, en in schnittstellen_werte:
        woerter.setdefault(de, en)
    synonyme: dict[str, str] = {}
    if "Synonyme" in wb.sheetnames:
        for r in list(wb["Synonyme"].iter_rows(values_only=True))[1:]:
            fid, begriff = _text(r[0]), _text(r[1] if len(r) > 1 else "")
            if fid and begriff and (fid in by_id):
                synonyme[normalisieren(begriff)] = fid
    # Synonyme aus dem mitgelieferten Katalog ergänzen (neue Programmversion → neue Begriffe,
    # auch wenn neben der EXE noch ein älterer, vom Nutzer gepflegter Katalog liegt).
    for begriff, fid in _mitgelieferte_synonyme(pfad).items():
        if fid in by_id:
            synonyme.setdefault(begriff, fid)
    for f in felder:                      # Feldnamen selbst gelten immer als Synonyme
        if f.abschnitt != INTERN:
            synonyme.setdefault(normalisieren(f.de), f.id)
            synonyme.setdefault(normalisieren(f.en), f.id)
    wb.close()
    formulierungen = _formulierungen(pfad)
    bezeichnungen = _bezeichnungen(pfad)
    return Katalog(felder=felder, woerterbuch=woerter, schnittstellen_werte=schnittstellen_werte, quelle=pfad,
                   synonyme=synonyme, formulierungen=formulierungen, bezeichnungen=bezeichnungen)


def _bezeichnungen_aus(pfad: Path) -> dict | None:
    try:
        wb = load_workbook(pfad, read_only=True, data_only=True)
    except Exception:
        return None
    if "Bezeichnungen" not in wb.sheetnames:
        wb.close()
        return None
    out = {}
    for r in list(wb["Bezeichnungen"].iter_rows(values_only=True))[1:]:
        r = list(r) + [None] * 4
        typ, fid, de, en = (_text(x) for x in r[:4])
        if typ and fid and de:
            out[(typ, fid)] = (de, en)
    wb.close()
    return out


def _bezeichnungen(pfad: Path) -> dict:
    """Blatt „Bezeichnungen“ aus dem Katalog im Programmordner; fehlt es dort (älterer Katalog), das mitgelieferte."""
    eigene = _bezeichnungen_aus(pfad)
    if eigene is not None:
        return eigene
    m = pfade.KATALOG_MITGELIEFERT
    try:
        if m.exists() and m.resolve() != Path(pfad).resolve():
            return _bezeichnungen_aus(m) or {}
    except OSError:
        pass
    return {}


def _formulierungen_aus(pfad: Path) -> list[tuple[str, str]]:
    try:
        wb = load_workbook(pfad, read_only=True, data_only=True)
    except Exception:
        return []
    out = []
    if "Formulierungen" in wb.sheetnames:
        for r in list(wb["Formulierungen"].iter_rows(values_only=True))[1:]:
            en, de = _text(r[0] if r else ""), _text(r[1] if r and len(r) > 1 else "")
            if en and de:
                out.append((en, de))
    wb.close()
    return out


def _formulierungen(pfad: Path) -> list[tuple[str, str]]:
    """Eigene Einträge aus dem Katalog im Programmordner + neue aus dem mitgelieferten Katalog."""
    eigene = _formulierungen_aus(pfad)
    bekannt = {en.lower() for en, _ in eigene}
    m = pfade.KATALOG_MITGELIEFERT
    try:
        if m.exists() and m.resolve() != Path(pfad).resolve():
            eigene += [(en, de) for en, de in _formulierungen_aus(m) if en.lower() not in bekannt]
    except OSError:
        pass
    return eigene


def _mitgelieferte_synonyme(pfad: Path) -> dict[str, str]:
    m = pfade.KATALOG_MITGELIEFERT
    try:
        if not m.exists() or m.resolve() == Path(pfad).resolve():
            return {}
        wb = load_workbook(m, read_only=True, data_only=True)
    except Exception:
        return {}
    erg: dict[str, str] = {}
    if "Synonyme" in wb.sheetnames:
        for r in list(wb["Synonyme"].iter_rows(values_only=True))[1:]:
            fid, begriff = _text(r[0]), _text(r[1] if len(r) > 1 else "")
            if fid and begriff:
                erg[normalisieren(begriff)] = fid
    wb.close()
    return erg


def normalisieren(s: str) -> str:
    s = (s or "").lower().strip()
    for z in ("■", "•", "–", "—"):
        s = s.replace(z, " ")
    s = s.strip(" -:.\t")
    s = re.sub(r"\s*/\s*", "/", s)
    return " ".join(s.split())

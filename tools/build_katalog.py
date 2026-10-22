"""Erzeugt katalog/feldkatalog.xlsx (Erst-Erstellung).

ACHTUNG: Die Excel-Datei ist die maßgebliche Quelle für den Converter und darf von Hand
gepflegt werden. Dieses Skript nur erneut ausführen, wenn der Katalog bewusst komplett
neu aufgebaut werden soll – es überschreibt die Datei.

Quellen:
  F = Filter-Auswahl Produktkonfigurator (22.09.2026, 55 Werte)
  V = bisherige EPS-Vorlage (kategorien.py DE 69 / kategorien_en.py EN 79)
  P = EPS-Datenblattpaare d/e 200034, 200161, 200326 (SelectLine)
  E = Entscheidungen aus dem Pflichtenheft (15.10.2026)
"""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "katalog" / "feldkatalog.xlsx"

# Datentypen (für Filter / Werte-Umwandlung); der Originaltext fürs PDF wird immer zusätzlich gespeichert.
T_TEXT, T_ZAHL, T_BEREICH, T_JN, T_AUSW, T_MULTI, T_LISTE, T_IF = (
    "Text", "Zahl", "Bereich (von–bis)", "Ja/Nein", "Auswahl", "Mehrfachauswahl", "Liste", "Schnittstelle")
OK, VOR = "bestätigt", "Vorschlag – bitte prüfen"

# (id, abschnitt, feld_de, feld_en, typ, einheit, kern, filterart, pflicht, quelle, status, bemerkung)
F = []


def f(id_, abschnitt, de, en, typ, einheit="", kern=False, filterart="", pflicht=False, quelle="", status=OK, bem=""):
    F.append((id_, abschnitt, de, en, typ, einheit, "ja" if kern else "", filterart, "ja" if pflicht else "",
              quelle, status, bem))


KOPF, ALLG, SCHN, TECH, OPT, LIEF, INT = ("Kopfdaten", "Allgemeine Daten", "Schnittstellen", "Technische Daten",
                                          "Optionen", "Lieferumfang", "Intern (nicht auf Datenblatt)")

# ---------------- Kopfdaten
f("geraetetyp", KOPF, "Gerätetyp", "Device type", T_MULTI, filterart="Einstiegsfrage", pflicht=True, quelle="E",
  bem="= Warengruppe in Weclapp. Bidirektionale Geräte zusätzlich „Elektronische Last“ (abwählbar).")
f("serie", KOPF, "Serie", "Series", T_TEXT, quelle="P")
f("typbezeichnung", KOPF, "Typbezeichnung", "Type designation", T_TEXT, pflicht=True, quelle="P, E",
  bem="EPS-Name: Original beginnt mit „PS“ → „E/“ + Original, sonst „EPS/“ + Original. Händlerware: Originalname.")
f("produkttyp", KOPF, "Produkttyp", "Product type", T_TEXT, quelle="P", bem="z. B. Labornetzgerät / Laboratory Power Supply")
f("beschreibung", KOPF, "Beschreibung", "Description", T_TEXT, quelle="P",
  bem="Freitext DE + EN; Quelle (übersetzter Herstellertext oder eigener EPS-Text) klärt der Chef (Punkt 3).")

# ---------------- Allgemeine Daten (Reihenfolge wie bisherige Vorlage, ergänzt)
f("technologie", ALLG, "Technologie", "Technology", T_AUSW, filterart="", quelle="V, F")
f("betriebsarten", ALLG, "Betriebsarten", "Operation modes", T_MULTI, kern=True, filterart="Mehrfachauswahl", quelle="V, F")
f("netzanschluss", ALLG, "Netzanschluss", "Mains", T_TEXT, quelle="V",
  bem="Text wie im Datenblatt (z. B. „208–240 V AC ±10 % 1ph.“). Die Eingangsspannung wird daraus abgeleitet.")
f("eingangsspannung", TECH, "Eingangsspannung", "Input Voltage", T_BEREICH, "V", kern=True, filterart="Schieberegler", quelle="F, V(EN)",
  bem="Für den Filter; Wert aus Netzanschluss bzw. Eingangsangaben des Herstellers.")
f("eingangsfrequenz", ALLG, "Eingangsfrequenz", "Input frequency", T_TEXT, "Hz", quelle="V")
f("leistungsfaktor", ALLG, "Leistungsfaktor", "Power factor", T_TEXT, quelle="V")
f("netzrueckspeisung", ALLG, "Netzrückspeisung", "Power feed back", T_JN, kern=True, filterart="Ja/Nein", quelle="V, F")
f("bidirektional", ALLG, "Bidirektional", "Bidirectional", T_JN, kern=True, filterart="Ja/Nein", quelle="F, P",
  status=VOR, bem="Im englischen Datenblatt bisher als „Behavior: Bidirectional“ (200326); im deutschen fehlte die Zeile.")
f("anzeige", ALLG, "Anzeige", "Display", T_AUSW, kern=True, filterart="Dropdown", quelle="V, F")
f("spannungsaufloesung", ALLG, "Spannungsauflösung", "Voltage resolution", T_TEXT, "V", quelle="V, F")
f("spannungsgenauigkeit", ALLG, "Spannungsgenauigkeit", "Voltage accuracy", T_TEXT, "%", quelle="V, F")
f("spannungsstabilitaet_last", ALLG, "Spannungsstabilität Last", "Voltage stability load", T_TEXT, "%", quelle="V, F")
f("spannungsstabilitaet_netz", ALLG, "Spannungsstabilität Netz", "Voltage stability mains", T_TEXT, "%", quelle="V, F")
f("spannungsausregelung_last", ALLG, "Spannungsausregelung Last", "Response time voltage", T_TEXT, "ms", quelle="V, F, P")
f("anstiegszeit_spannung", ALLG, "Anstiegszeit Spannung", "Rise time voltage", T_TEXT, "ms", quelle="V, F")
f("stromaufloesung", ALLG, "Stromauflösung", "Current resolution", T_TEXT, "A", quelle="V, F")
f("stromgenauigkeit", ALLG, "Stromgenauigkeit", "Current accuracy", T_TEXT, "%", quelle="V, F")
f("stromstabilitaet_last", ALLG, "Stromstabilität Last", "Current stability load", T_TEXT, "%", quelle="V, F")
f("stromstabilitaet_netz", ALLG, "Stromstabilität Netz", "Current stability mains", T_TEXT, "%", quelle="V, F")
f("stromausregelung", ALLG, "Stromausregelung", "Response time current", T_TEXT, "ms", quelle="F", status=VOR,
  bem="Nur in der Filter-Auswahl; englische Bezeichnung neu vorgeschlagen.")
f("anstiegszeit_strom", ALLG, "Anstiegszeit Strom", "Rise time current", T_TEXT, "ms", quelle="V, F")
f("begrenzung_ausgangsstrom", ALLG, "Begrenzung Ausgangsstrom", "Output current limitation", T_JN, quelle="V, F")
f("eingangsstrombegrenzung", ALLG, "Eingangsstrombegrenzung", "Input current limitation", T_JN, quelle="V(EN)", status=VOR,
  bem="Bisher nur in der englischen Vorlage. Prüfen: identisch mit „Einschaltstrombegrenzung“?")
f("einschaltstrombegrenzung", ALLG, "Einschaltstrombegrenzung", "Inrush current limitation", T_JN, quelle="F", status=VOR,
  bem="Nur in der Filter-Auswahl; englische Bezeichnung neu vorgeschlagen.")
f("leistungsgenauigkeit", ALLG, "Leistungsgenauigkeit", "Power accuracy", T_TEXT, "%", quelle="V, F")
f("innenwiderstandsregelung", ALLG, "Innenwiderstandsregelung", "Internal resistance regulation", T_JN, quelle="V, F",
  bem="Ja/Nein für den Filter, Zusatztext (z. B. „+Simulation“) bleibt fürs Datenblatt.")
f("ueberspannungskategorie", ALLG, "Überspannungskategorie", "Overvoltage category", T_AUSW, quelle="V, F")
f("ueberhitzungsschutz", ALLG, "Überhitzungsschutz", "Overheat protection", T_JN, quelle="V, F")
f("spannungsfestigkeit_ein_aus", ALLG, "Spannungsfestigkeit Eingang zu Ausgang", "Isolation input/output", T_TEXT, "V", quelle="V")
f("spannungsfestigkeit_aus_geh", ALLG, "Spannungsfestigkeit Ausgang zu Gehäuse", "Isolation output/enclosure", T_TEXT, "V", quelle="V")
f("schutzklasse", ALLG, "Schutzklasse", "Protection class", T_AUSW, quelle="V, F")
f("reihenschaltung", ALLG, "Reihenschaltung", "Series operation", T_JN, quelle="V", status=VOR,
  bem="Englische Zeile fehlte bisher. Bis zur Klärung (Chef-Punkt 2) getrennt von Master-Slave Reihenschaltung.")
f("master_slave_reihenschaltung", ALLG, "Master-Slave Reihenschaltung", "Master-slave series operation", T_JN, quelle="F", status=VOR,
  bem="Chef-Punkt 2: identisch mit „Reihenschaltung“?")
f("parallelschaltung", ALLG, "Parallelschaltung", "Parallel operation", T_AUSW, kern=True, filterart="Dropdown", quelle="V, F")
f("current_sharing", ALLG, "Current Sharing", "Current sharing", T_JN, quelle="V, F")
f("kuehlung", ALLG, "Kühlung", "Cooling", T_AUSW, kern=True, filterart="Dropdown", quelle="V, F")
f("betriebstemperatur", ALLG, "Betriebstemperatur", "Operating temperature", T_BEREICH, "°C", kern=True, filterart="Schieberegler", quelle="V, F",
  bem="Bisher EN „Operation temperature“.")
f("lagertemperatur", ALLG, "Lagertemperatur", "Storage temperature", T_BEREICH, "°C", quelle="V, F")
f("luftfeuchtigkeit", ALLG, "Luftfeuchtigkeit", "Humidity", T_TEXT, "%", quelle="V, F")
f("betriebshoehe", ALLG, "Betriebshöhe", "Altitude", T_TEXT, "m", quelle="V, F", bem="Bisher falsch übersetzt als „Attitude“.")
f("bauform", ALLG, "Bauform", "Design", T_AUSW, kern=True, filterart="Dropdown", quelle="V, F")
f("normen", ALLG, "Normen", "Standards", T_TEXT, quelle="V, F")
f("robust", ALLG, "Robuste Ausführung", "Ruggedized", T_JN, quelle="V(EN)", status=VOR, bem="Bisher nur in der englischen Vorlage („ruggedized“).")
f("power_fail", ALLG, "Power fail", "Power fail", T_JN, quelle="V, F")
f("alarmmanagement", ALLG, "Alarmmanagement", "Alarm management", T_JN, quelle="V, F", bem="Bisher EN „Alarmmanagement“ (deutsch).")
f("funktionsgenerator", ALLG, "Funktionsgenerator", "Function generator", T_JN, kern=True, filterart="Ja/Nein", quelle="V, F",
  bem="Ja/Nein für den Filter, Zusatztext (z. B. „+arbitrary“) bleibt fürs Datenblatt.")
f("voreinstellung_ausgang", ALLG, "Voreinstellung Ausgang", "Output preset", T_JN, quelle="V, F")
f("speicherplaetze", ALLG, "Speicherplätze", "Memory", T_TEXT, quelle="V, F")
f("innenkapazitaet", ALLG, "Kapazität", "Capacity", T_TEXT, "µF", quelle="V")

# ---------------- Schnittstellen (je Zeile: Standard / Option + Code; Filter „Schnittstellen“ wird daraus abgeleitet)
f("schnittstellen", SCHN, "Schnittstellen", "Interfaces", T_MULTI, kern=True, filterart="Mehrfachauswahl", quelle="F",
  bem="Wird automatisch aus den Schnittstellen-Zeilen unten gebildet (Standard oder Option).")
for id_, de, en, q, st, bem in [
    ("if_analog", "Analoge Programmierung (isoliert)", "Analog programming (isolated)", "V, P", VOR, "Bisher DE „Analoge Programmierung ISO“, EN „Analog Isolation“."),
    ("if_usb", "USB-Schnittstelle", "USB interface", "V, P", OK, ""),
    ("if_ethernet", "Ethernet-Schnittstelle", "Ethernet interface", "V, P", OK, ""),
    ("if_rs232", "RS232-Schnittstelle", "RS232 interface", "V, P", OK, ""),
    ("if_rs485", "RS485-Schnittstelle", "RS485 interface", "V", OK, ""),
    ("if_gpib", "GPIB-Schnittstelle", "GPIB interface", "V", OK, "Bisher falsch „GBIP“."),
    ("if_can", "CAN-Schnittstelle", "CAN interface", "V, P", OK, ""),
    ("if_canopen", "CANopen-Schnittstelle", "CANopen interface", "Beispiele", VOR, "Neu: in 5 Hersteller-PDFs genannt."),
    ("if_profibus", "Profibus", "Profibus", "V, P", OK, ""),
    ("if_profinet", "Profinet", "Profinet", "P(EN)", VOR, "Bisher nur im englischen Datenblatt."),
    ("if_modbus", "Modbus", "Modbus", "P(EN)", VOR, "Bisher nur im englischen Datenblatt."),
    ("if_ethercat", "EtherCAT-Schnittstelle", "EtherCAT interface", "V, P", OK, ""),
]:
    f(id_, SCHN, de, en, T_IF, quelle=q, status=st, bem=bem)
f("genauigkeit_schnittstelle", SCHN, "Genauigkeit Schnittstelle", "Interface accuracy", T_TEXT, quelle="V, P")
f("software", SCHN, "Software", "Software", T_TEXT, quelle="V, P")

# ---------------- Technische Daten
f("ausgangsspannung", TECH, "Ausgangsspannung", "Output voltage", T_BEREICH, "V", kern=True, filterart="Schieberegler", pflicht=True, quelle="V, F",
  bem="Filter: Wunsch +20 %, aufgerundet (bis 100: 1-2-5-Reihe, darüber 250er-Schritte), harte Obergrenze.")
f("ausgangsspannung_2", TECH, "Ausgangsspannung 2", "Output voltage 2", T_BEREICH, "V", quelle="V(EN)", status=VOR,
  bem="Bisher nur in der englischen Vorlage (Geräte mit zweitem Ausgang/Bereich).")
f("ausgangsstrom", TECH, "Ausgangsstrom", "Output current", T_BEREICH, "A", kern=True, filterart="Schieberegler", pflicht=True, quelle="V, F")
f("ausgangsstrom_2", TECH, "Ausgangsstrom 2", "Output current 2", T_BEREICH, "A", quelle="V(EN)", status=VOR)
f("ausgangsleistung", TECH, "Ausgangsleistung", "Output power", T_BEREICH, "W", kern=True, filterart="Schieberegler", pflicht=True, quelle="V, F")
f("eingangsstrom", TECH, "Eingangsstrom", "Input current", T_BEREICH, "A", quelle="F, P")
f("eingangsleistung", TECH, "Eingangsleistung", "Input power", T_BEREICH, "W", quelle="F", status=VOR, bem="Nur in der Filter-Auswahl.")
f("wirkungsgrad", TECH, "Wirkungsgrad", "Efficiency", T_TEXT, "%", quelle="V, F")
f("restwelligkeit_u", TECH, "Restwelligkeit U", "Ripple U", T_TEXT, quelle="V, F")
f("restwelligkeit_i", TECH, "Restwelligkeit I", "Ripple I", T_TEXT, quelle="V, F")
f("widerstand_einstellbereich", TECH, "Widerstand Einstellbereich", "Resistance adjustment range", T_BEREICH, "Ω", kern=True, filterart="Schieberegler", quelle="V, F",
  bem="Bisher „Widerstand Einstellbereich 1“.")
f("widerstand_aufloesung", TECH, "Widerstand Auflösung", "Resistance resolution", T_TEXT, quelle="V, P")
f("fernfuehlung", TECH, "Fernfühlungsausregelung", "Remote sensing", T_JN, quelle="V, P")
f("ausgangsfrequenz", TECH, "Ausgangsfrequenz", "Output frequency", T_AUSW, "Hz", kern=True, filterart="Dropdown", quelle="V, F")
f("frequenzgenauigkeit", TECH, "Frequenzgenauigkeit", "Frequency accuracy", T_TEXT, quelle="V")
f("scheitelfaktor", TECH, "Scheitelfaktor", "Crest factor", T_TEXT, quelle="V")
f("klirrfaktor", TECH, "Klirrfaktor (THD)", "Distortion (THD)", T_TEXT, "%", quelle="V, F")
f("abmessungen", TECH, "Abmessungen (B × H × T)", "Dimensions (W × H × D)", T_TEXT, "mm", quelle="V, F")
f("gewicht", TECH, "Gewicht", "Weight", T_ZAHL, "kg", quelle="V, F", bem="Geht zusätzlich als Artikelgewicht nach Weclapp.")
f("bestellnummer", TECH, "Bestellnummer", "Order code", T_TEXT, kern=True, filterart="SKU-Suche", pflicht=True, quelle="V, F, E",
  bem="Eingabe von Hand; Dateiname d/e + Bestellnummer.")

# ---------------- Listen
f("optionen", OPT, "Optionen", "Options", T_LISTE, kern=True, filterart="Suchtextfeld (Optionsnummer)", quelle="V, F, E",
  bem="Je Eintrag: Optionsnummer (einheitlich, Stammliste folgt in Weclapp), EPS-Code, Beschreibung DE/EN.")
f("lieferumfang", LIEF, "Lieferumfang", "Scope of delivery", T_LISTE, quelle="V, P")

# ---------------- Intern
f("hersteller", INT, "Hersteller", "Manufacturer", T_TEXT, pflicht=True, quelle="E", bem="Weclapp-Lieferantendaten; nie auf dem EPS-Datenblatt.")
f("hersteller_typ", INT, "Hersteller-Typbezeichnung", "Manufacturer type", T_TEXT, quelle="E")
f("hersteller_bestellnummer", INT, "Hersteller-Bestellnummer", "Manufacturer order code", T_TEXT, quelle="E",
  bem="Wird beim Import auf Doppelte geprüft (Hinweis).")
f("haendlerware", INT, "Händlerware (Originalname + Original-PDF)", "Reseller product", T_JN, quelle="E",
  bem="Schalter pro Hersteller, vorbereitet für EA/Tek; aus, bis J.-D. Enkelmann freischaltet.")
f("herkunft", INT, "Herkunft", "Origin", T_AUSW, quelle="E", bem="Neu (Auftrag) / Altbestand")
f("shop_sichtbar", INT, "Im Shop sichtbar", "Visible in shop", T_JN, quelle="E", bem="Voreinstellung: Altbestand ja, neue Auftragsgeräte nein.")
f("bilder", INT, "Produktbilder", "Product images", T_LISTE, quelle="E", bem="Von Hand hinzufügbar; Herkunft klärt der Chef (Punkt 6).")

# ---------------- Auswahlwerte: (feld_id, wert_de, wert_en, status, beleg)
A = []
def a(fid, de, en, st=VOR, beleg=""):
    A.append((fid, de, en, st, beleg))

for de, en in [("AC/AC", "AC/AC"), ("AC/DC", "AC/DC"), ("DC/AC", "DC/AC"), ("DC/DC", "DC/DC"),
               ("Elektronische Last", "Electronic load"), ("Ladegeräte", "Chargers")]:
    a("geraetetyp", de, en, OK, "Pflichtenheft (= Warengruppen)")
a("technologie", "Switching", "Switching", VOR, "EPS-Datenblätter (alle 3 Paare)")
a("technologie", "Linear", "Linear", VOR, "Vorschlag")
for v in ["CV", "CC", "CP", "CR"]:
    a("betriebsarten", v, v, VOR, "EPS-Datenblätter 200034/200161/200326")
for de, en, b in [("TFT-Touch-Display", "TFT touch display", "EPS 200034, 200161"),
                  ("TFT-Display", "TFT display", "EPS 200326"),
                  ("LCD", "LCD", "Hersteller-PDFs (häufig)"),
                  ("LED", "LED", "Hersteller-PDFs"),
                  ("ohne Anzeige (nur Fernsteuerung)", "None (remote control only)", "Vorschlag")]:
    a("anzeige", de, en, VOR, b)
for de, en, b in [("Nein", "No", "Vorschlag"), ("Ja", "Yes", "EPS 200161 („Standard“)"),
                  ("Master-Slave", "Master-slave", "EPS 200034, 200326; Plating")]:
    a("parallelschaltung", de, en, VOR, b)
for de, en, b in [("Lüfter", "Fan", "EPS-Datenblätter, Absopulse („Air Forced“)"),
                  ("Wasser", "Water", "Gustav Klein, EPS-Option EPS/WC"),
                  ("Konvektion", "Convection", "Vorschlag")]:
    a("kuehlung", de, en, VOR, b)
for de, en, b in [("19 Zoll", "19 inch", "EPS 200034, 200326; Absopulse, M.C.B"),
                  ("Desktop", "Desktop", "EPS 200161"),
                  ("Tower", "Tower", "M.C.B („Mini Tower“)"),
                  ("Schrank", "Cabinet", "Gustav Klein, Plating"),
                  ("Einbaumodul", "Module", "Vorschlag")]:
    a("bauform", de, en, VOR, b)
for de, en, b in [("DC", "DC", "alle DC-Geräte"), ("50 Hz", "50 Hz", "SHT, Gustav Klein"), ("60 Hz", "60 Hz", "SHT"),
                  ("50/60 Hz", "50/60 Hz", "SHT, Gustav Klein"), ("400 Hz", "400 Hz", "Absopulse (FTP 400 Hz)"),
                  ("einstellbar", "adjustable", "SHT („40–70 Hz“)")]:
    a("ausgangsfrequenz", de, en, VOR, b)
for v in ["1", "2", "3", "4"]:
    a("ueberspannungskategorie", v, v, VOR, "EPS-Datenblätter („2“)")
for v in ["1", "2", "3"]:
    a("schutzklasse", v, v, VOR, "EPS-Datenblätter („1“)")
for de, en in [("Neu (Auftrag)", "New (order)"), ("Altbestand", "Existing stock")]:
    a("herkunft", de, en, OK, "Pflichtenheft")
for de, en in [("Standard", "Standard"), ("Option", "Option"), ("—", "—")]:
    a("Schnittstelle (alle if_*)", de, en, VOR, "EPS-Datenblätter; bei „Option“ zusätzlich EPS-Optionscode")

# ---------------- Wörterbuch (Werte und Texte, aus den EPS-Paaren; EN teils korrigiert)
W = []
def w(de, en, quelle, bem=""):
    W.append((de, en, quelle, bem))

for de, en, q, b in [
    ("Labornetzgerät", "Laboratory power supply", "200034, 200161", ""),
    ("Bidirektionales DC-Netzgerät + Netzrückspeisung", "Bidirectional DC power supply + energy recovery", "200326", ""),
    ("Netzgerät", "Power supply unit", "200034, 200161", ""),
    ("Bidirektionales Netzgerät", "Bidirectional power supply unit", "200326", "DE-Datenblatt sagte nur „Netzgerät“"),
    ("Testprotokoll", "Test report", "alle", ""),
    ("2x Stecker für Fernfühlung", "2x plug for remote sense", "200034, 200326", ""),
    ("Set DC-Klemmenabdeckung", "Set DC terminal cover", "200034, 200326", ""),
    ("Set Klemmenabdeckung Sense", "Set terminal cover sense", "200034, 200326", "EN bisher „Terminal cover sense“"),
    ("Kabelbinder für Zugentlastung", "Cable tie for strain relief", "200034", ""),
    ("USB-Kabel 1,8 m", "USB cable 1.8 m", "alle", ""),
    ("USB-Stick mit Dokumentation und Software", "USB stick with documentation and software", "200034, 200326", ""),
    ("USB-Stick mit Handbuch und Software", "USB stick with manual and software", "200161", ""),
    ("Netzkabel 1,5 bzw. 2,0 m", "Mains cord 1.5 or 2.0 m", "200161", ""),
    ("AC-Anschlussstecker (Klemmtyp)", "AC connector plug (clamp type)", "200326", ""),
    ("Set AC-Kabel-Zugentlastung", "Set for AC cable strain relief", "200326", ""),
    ("Kalibrierung mit Protokoll", "Calibration with protocol", "alle", "EN bisher „protocoll“ (Tippfehler)"),
    ("Garantieverlängerung", "Extended warranty", "200034, 200326", ""),
    ("Wasserkühlung", "Water cooling", "200326", "EN bisher „Watercooling“"),
    ("19\"-Einbaurahmen", "19\" mounting frame", "200161", ""),
    ("Modbus-Schnittstelle 1/2 Port", "Modbus interface 1/2 port", "200034", ""),
    ("Netz- und Systemüberwachungseinheit", "Grid and system protection unit", "200326", ""),
    ("Share-Bus-Verbindungskabel", "Share bus connection cable", "200034, 200326", ""),
    ("Master-Slave-Patchkabel", "Master-slave patch cable", "200034, 200326", ""),
    ("Lüfter", "Fan", "alle", ""),
    ("19 Zoll", "19 inch", "200034, 200326", ""),
    ("Standard", "Standard", "alle", ""),
    ("Option", "Option", "alle", ""),
    ("Derating", "derating", "alle", ""),
    ("Irrtümer und Änderungen vorbehalten / Alle Wertangaben sind typische Werte",
     "Errors and changes excepted / All values are typical values", "alle", "Fußzeile"),
]:
    w(de, en, q, b)

# ---------------- Excel schreiben
ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="35528C")
KERN_FILL = PatternFill("solid", fgColor="E3F1F8")
VOR_FILL = PatternFill("solid", fgColor="FFF2B3")
thin = Side(style="thin", color="D5DCE6")
BORDER = Border(bottom=thin)


def sheet(ws, headers, rows, widths, kern_col=None, status_col=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=ARIAL, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
        if kern_col is not None and row[kern_col].value == "ja":
            for c in row[:kern_col + 1]:
                c.fill = KERN_FILL
        if status_col is not None and row[status_col].value == VOR:
            row[status_col].fill = VOR_FILL
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


wb = Workbook()
ws0 = wb.active
ws0.title = "Anleitung"
lines = [
    ("Feldkatalog EPS-Datenblatt-Converter", True),
    ("Stand: 22.10.2026 · Grundlage: Pflichtenheft Produktdaten-Pipeline (15.10.2026)", False),
    ("", False),
    ("Wofür ist diese Datei?", True),
    ("Der Converter liest diese Datei beim Start ein. Sie legt fest, welche Felder es gibt, wie sie auf Deutsch und Englisch heißen,", False),
    ("welcher Datentyp gilt, welche Felder Kernfelder (Konfigurator-Filter) sind und welche Auswahlwerte erlaubt sind.", False),
    ("", False),
    ("Was darfst du ändern?", True),
    ("• Spalten „Feld (DE)“, „Feld (EN)“, „Bemerkung“ und die Reihenfolge der Zeilen (= Reihenfolge auf dem Datenblatt).", False),
    ("• Blatt „Auswahlwerte“: Werte ergänzen, umbenennen, löschen. Danach Status auf „bestätigt“ setzen.", False),
    ("• Blatt „Wörterbuch“: Übersetzungen ergänzen oder korrigieren.", False),
    ("• Spalte „ID“ bitte NICHT ändern – darüber werden die Daten gespeichert.", False),
    ("", False),
    ("Farben", True),
    ("Hellblau = eines der 18 Kernfelder (müssen in Weclapp in jedem Fall stimmen, Pflichtprüfung in der Prüfmaske).", False),
    ("Gelb (Spalte Status) = Vorschlag von Claude, bitte prüfen und auf „bestätigt“ setzen.", False),
    ("", False),
    ("Quellen (Spalte „Quelle“)", True),
    ("F = Filter-Auswahl Produktkonfigurator · V = bisherige EPS-Vorlage (DE/EN) · P = EPS-Datenblattpaare 200034/200161/200326 · E = Pflichtenheft", False),
    ("", False),
    ("Beispiel", True),
    ("Feld „Ausgangsspannung“: Datentyp „Bereich (von–bis)“, Einheit V. Im Datenblatt steht „0–750 V DC“; gespeichert wird zusätzlich min 0 / max 750 / V für den Filter.", False),
]
for txt, bold in lines:
    ws0.append([txt])
    ws0.cell(ws0.max_row, 1).font = Font(name=ARIAL, bold=bold, size=12 if bold and ws0.max_row == 1 else 10,
                                         color="35528C" if bold else "000000")
ws0.column_dimensions["A"].width = 140

ws1 = wb.create_sheet("Felder")
hdr = ["ID", "Abschnitt", "Feld (DE)", "Feld (EN)", "Datentyp", "Einheit", "Kernfeld (18)", "Filterart Konfigurator",
       "Pflicht", "Quelle", "Status", "Bemerkung"]
sheet(ws1, hdr, F, [26, 22, 34, 32, 18, 8, 10, 22, 8, 12, 22, 70], kern_col=6, status_col=10)
dv_typ = DataValidation(type="list", formula1='"' + ",".join([T_TEXT, T_ZAHL, T_BEREICH, T_JN, T_AUSW, T_MULTI, T_LISTE, T_IF]) + '"', allow_blank=True)
dv_jn = DataValidation(type="list", formula1='"ja,"', allow_blank=True)
dv_st = DataValidation(type="list", formula1=f'"{OK},{VOR}"', allow_blank=True)
for dv in (dv_typ, dv_jn, dv_st):
    ws1.add_data_validation(dv)
n = ws1.max_row
dv_typ.add(f"E2:E{n + 200}"); dv_jn.add(f"G2:G{n + 200}"); dv_jn.add(f"I2:I{n + 200}"); dv_st.add(f"K2:K{n + 200}")

ws2 = wb.create_sheet("Auswahlwerte")
sheet(ws2, ["Feld-ID", "Wert (DE)", "Wert (EN)", "Status", "Beleg"], A, [28, 34, 34, 22, 50], status_col=3)
dv_st2 = DataValidation(type="list", formula1=f'"{OK},{VOR}"', allow_blank=True)
ws2.add_data_validation(dv_st2); dv_st2.add(f"D2:D{ws2.max_row + 200}")

ws3 = wb.create_sheet("Wörterbuch")
sheet(ws3, ["Deutsch", "Englisch", "Quelle (EPS-Artikel)", "Bemerkung"], W, [50, 50, 20, 45])

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
kern = sum(1 for r in F if r[6] == "ja")
print(f"{OUT}: {len(F)} Felder, davon {kern} Kernfelder; {len(A)} Auswahlwerte; {len(W)} Wörterbuch-Einträge")

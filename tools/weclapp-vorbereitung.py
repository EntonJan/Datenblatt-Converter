#!/usr/bin/env python3
"""Vorlagen für die Weclapp-Einführung erzeugen.

    python tools/weclapp-vorbereitung.py <shopdaten.json> <Zielordner>

Erzeugt:
  Weclapp-Feldliste.xlsx      – eigene Artikelattribute (Schlüssel, Bezeichnung DE/EN, Typ, Auswahlwerte) + Standardfelder
  Weclapp-Artikelimport.xlsx  – alle Shopgeräte im Importformat (zum Ansehen/Prüfen)
  Weclapp-Artikelimport.csv   – dieselben Daten für den Weclapp-Import (UTF-8, Semikolon)
"""
import csv
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook                                   # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill        # noqa: E402
from openpyxl.utils import get_column_letter                    # noqa: E402

from eps_converter import katalog as K, pfade                  # noqa: E402
from eps_produktsuche import shop, weclapp_felder as WF         # noqa: E402

KOPF = PatternFill("solid", fgColor="35528C")
HELL = PatternFill("solid", fgColor="EEF4FA")
GELB = PatternFill("solid", fgColor="FFF4D6")


def blatt(wb, titel, kopf, zeilen, breiten, erste=False):
    ws = wb.active if erste else wb.create_sheet()
    ws.title = titel
    ws.append(kopf)
    for z in zeilen:
        ws.append(list(z))
    for i, c in enumerate(ws[1], 1):
        c.font, c.fill = Font(bold=True, color="FFFFFF"), KOPF
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = breiten[i - 1] if i <= len(breiten) else 14
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for zeile in ws.iter_rows(min_row=2):
        for c in zeile:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def main(quelle: str, ziel: str) -> None:
    ziel = Path(ziel)
    ziel.mkdir(parents=True, exist_ok=True)
    kat = K.laden(pfade.KATALOG_MITGELIEFERT)
    attribute, standard = WF.feldliste(kat)

    wb = Workbook()
    ws = blatt(wb, "Anleitung", ["Weclapp-Feldliste – Vorlage für die Einrichtung"], [
        [f"Stand {date.today():%d.%m.%Y} · Grundlage: Feldkatalog des Datenblatt-Converters + Felder der Produktsuche"],
        [""],
        ["1. Blatt „Eigene Attribute“: in Weclapp unter Einstellungen → Eigene Attribute → Artikel anlegen."],
        ["   Schlüssel (eps_…) genau so übernehmen – Produktsuche, Konfigurator und Converter lesen über diese Schlüssel."],
        ["   Typ: STRING = Text, LARGE_TEXT = langer Text, DECIMAL = Zahl, INTEGER = ganze Zahl, BOOLEAN = Ja/Nein,"],
        ["   LIST = Auswahl (eine), MULTISELECT_LIST = Mehrfachauswahl, URL = Link. Auswahlwerte stehen in Spalte F."],
        ["2. Blatt „Standardfelder“: diese Werte haben in Weclapp eigene Felder (Artikelnummer, Name, Lieferant …)."],
        ["3. Bereiche (von–bis) sind dreifach angelegt: Text fürs Datenblatt („0-60 VDC“) + min + max als Zahl"],
        ["   (Grundeinheit V, A, W, °C, Ω) für Filter und Suche."],
        ["4. Gelb markiert = Kernfelder (Filter auf der Homepage / im Konfigurator)."],
        [""],
        ["Offen: Paket (ERP Handel), Artikelnummer = Shop-Bestellnummer?, Übersetzungen (DE/EN) in Weclapp,"],
        ["       wer richtet ein. Feldnamen in der API im 30-Tage-Test prüfen."],
    ], [120], erste=True)
    ws.freeze_panes = None
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    zeilen = [[a.schluessel, a.de, a.en, a.typ, a.einheit, a.auswahl_de, a.auswahl_en, "ja" if a.pflicht else "",
               a.filter, a.abschnitt, a.bemerkung] for a in attribute]
    ws = blatt(wb, "Eigene Attribute", ["Schlüssel", "Bezeichnung DE", "Bezeichnung EN", "Typ", "Einheit",
                                        "Auswahlwerte DE", "Auswahlwerte EN", "Pflicht", "Filter (Kernfeld)", "Abschnitt",
                                        "Bemerkung"], zeilen, [30, 28, 28, 18, 8, 40, 40, 8, 18, 22, 50])
    for zeile in ws.iter_rows(min_row=2):
        if zeile[8].value:
            for c in zeile:
                c.fill = GELB
    blatt(wb, "Standardfelder", ["Feld (Converter)", "Bezeichnung", "Weclapp-Feld", "Bemerkung"], standard,
          [24, 30, 40, 70])
    wb.save(ziel / "Weclapp-Feldliste.xlsx")

    geraete, stand = shop.laden(Path(quelle))
    kopf = [s for s, _ in WF.IMPORT_SPALTEN]
    zeilen = WF.artikel_zeilen(geraete)
    wb = Workbook()
    ws = blatt(wb, "Artikel", kopf, zeilen, [11, 40, 16, 7, 34, 20, 22] + [16] * (len(kopf) - 8) + [50], erste=True)
    for zeile in ws.iter_rows(min_row=2):
        if zeile[-1].value:
            zeile[-1].fill = GELB
        zeile[5].fill = zeile[6].fill = HELL
    blatt(wb, "Hinweise", ["Hinweise zur Importdatei"], [
        [f"Quelle: Shopdaten der EPS-Produktsuche (Stand {stand[:16].replace('T', ' ')}), {len(zeilen)} Geräte."],
        ["Lieferant und Hersteller-Typ stehen nicht im Shop – bitte ausfüllen (hellblau) oder später in Weclapp pflegen."],
        ["Spalte „Prüfen“ (gelb): bekannte Datenfehler aus dem Shop bzw. Werte nur aus der Typbezeichnung."],
        ["Ausführliche technische Werte kommen später aus dem Datenblatt-Converter (freigegebene Datensätze)."],
        ["Zahlen mit Komma als Dezimaltrennzeichen; Bereiche min/max in V, A, W."],
    ], [110])
    wb.save(ziel / "Weclapp-Artikelimport.xlsx")
    with open(ziel / "Weclapp-Artikelimport.csv", "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(kopf)
        w.writerows(zeilen)
    print(f"{len(attribute)} eigene Attribute, {len(standard)} Standardfelder, {len(zeilen)} Artikel → {ziel}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])

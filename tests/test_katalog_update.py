"""Neue Synonyme aus dem mitgelieferten Katalog wirken auch, wenn der Nutzer einen älteren Katalog hat."""
import shutil
from pathlib import Path

from openpyxl import load_workbook

from eps_converter import katalog, pfade

KAT = Path(__file__).resolve().parents[1] / "katalog" / "feldkatalog.xlsx"


def test_synonyme_aus_mitgeliefertem_katalog(tmp_path, monkeypatch):
    alt = tmp_path / "feldkatalog.xlsx"
    shutil.copy2(KAT, alt)
    wb = load_workbook(alt)
    ws = wb["Synonyme"]
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 2).value or "").lower().startswith("package"):
            ws.delete_rows(r)
    wb.save(alt)
    monkeypatch.setattr(pfade, "KATALOG_MITGELIEFERT", tmp_path / "fehlt.xlsx")
    assert katalog.normalisieren("Package / Dimensions") not in katalog.laden(alt).synonyme
    monkeypatch.setattr(pfade, "KATALOG_MITGELIEFERT", KAT)
    assert katalog.laden(alt).synonyme[katalog.normalisieren("Package / Dimensions (WxHxL)".split(" (")[0])] == "abmessungen"
